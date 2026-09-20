# -*- coding: utf-8 -*-
"""各输入框的上传模板。

设计上有一条硬规矩:**说明文字一律放单独的工作表**。
如果把说明写在数据列里,用户填完原样传回来,那些说明会被当成关键词导进去。
模板必须能安全地「下载 → 填 → 直接传回」。
"""
import io

TEMPLATES = {
    "seeds": {
        "title": "种子词",
        "columns": ["种子词"],
        "examples": [["conveyor roller"], ["idler roller"], ["gravity conveyor"]],
        "notes": [
            "一行一个种子词，第一行表头请保留。",
            "GKP 每次最多吃 20 个种子，超过会自动分批并去重，填多少都行。",
            "只读第一列，右边多出来的列会被忽略——直接拿别的表改也可以。",
        ],
    },
    "words": {
        "title": "关键词",
        "columns": ["关键词"],
        "examples": [["conveyor rollers"], ["gravity roller conveyor"],
                     ["heavy duty roller conveyor"]],
        "notes": [
            "一行一个关键词，一次最多 1 万个。",
            "也可以把关键词规划师网页版导出的 CSV 原样传上来——",
            "那种文件是 UTF-16 + Tab 分隔、前两行是标题和日期，程序会自动识别并剥掉。",
            "导入时按大小写不敏感去重。",
        ],
    },
    "sites": {
        "title": "竞品网址",
        "columns": ["竞品网址"],
        "examples": [["https://competitor-a.com"], ["https://competitor-b.com"],
                     ["competitor-c.com"]],
        "notes": [
            "一行一个，带不带 https:// 都行。",
            "按整站拓词，来源渠道列会记成「GKP以网站拓展(域名)」。",
            "网址取不到词时会跳过并在日志里说明，不影响其余的。",
        ],
    },
    "customer": {
        "title": "客户原始词",
        "columns": ["关键词", "客户中文", "客户级别"],
        "examples": [["conveyor roller", "输送机辊筒", "S"],
                     ["rubber roller", "橡胶辊筒", "A+"],
                     ["return roller", "改向滚筒", "A"]],
        "notes": [
            "第一列必填，后两列可以留空。",
            "客户级别按你们自己的分档填（S / A+ / A / B / C），程序不做校验，原样带走。",
            "这三列会原样进「2.原始词核对」页；这些词在总表里词源标成「原始词」。",
            "也可以在文本框里直接粘，用 Tab 或逗号分隔三段即可。",
        ],
    },
    "mixed": {
        "title": "混杂泛词清单",
        "columns": ["词 / 模式"],
        "examples": [["roller manufacturer"], ["rubber roller*"], ["steel roller*"]],
        "notes": [
            "命中的词，优先级最高只给 P1（对应口径说明第 8 条「混杂意图词最高只给 P1」）。",
            "",
            "匹配语法（和剔除清单同一套）：",
            "  roller manufacturer   整词——只命中这个词组本身",
            "  rubber roller*        前缀——命中 rubber roller、rubber roller price…",
            "  *conveyor belt*       包含——命中任何含这个词组的词",
            "",
            "默认整词是刻意的：roller manufacturer 要降级，但 conveyor roller manufacturer",
            "是带行业限定的合格词，用包含匹配会把后者一起误伤。",
        ],
    },
    "exclude": {
        "title": "剔除清单",
        "columns": ["词 / 模式"],
        "examples": [["*conveyor belt*"], ["*conveyor system*"], ["roller"], ["pulley"]],
        "notes": [
            "命中的词移出总表，单独进「4.剔除词」页，剔除原因记「命中剔除清单」。",
            "",
            "匹配语法和混杂清单相同：整词 / 前缀* / *包含*。",
            "",
            "对照口径说明的剔除规则该怎么写：",
            "  ①泛词（roller、pulley 单独成词）      → 写整词：roller",
            "  ②跨行业混杂（油漆辊、轮滑轮）         → 写包含：*paint roller*",
            "  ③系统级词（整机、系统）              → 写包含：*conveyor system*",
            "  ④输送带/链条/轴承类                  → 写包含：*conveyor belt*",
            "  ⑤不生产的产品线                      → 写包含：*drum motor*",
            "  ⑥品牌/平台词                        → 写包含：*amazon*",
            "  ⑦月搜过低                           → 不用写，界面上设「最低月搜」即可",
        ],
    },
}


def build(kind):
    """生成模板 xlsx 的字节流,返回 (bytes, 文件名)。"""
    spec = TEMPLATES.get(kind)
    if not spec:
        raise KeyError("没有这个模板:%s" % kind)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = spec["title"]
    ws.append(spec["columns"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4A5568")
    for row in spec["examples"]:
        ws.append(row)
    for i in range(len(spec["columns"])):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = \
            [30, 18, 12][i] if i < 3 else 20
    ws.freeze_panes = "A2"

    # 说明单独一页 —— 不能混进数据列,否则填完传回来会被当成关键词
    note = wb.create_sheet("怎么填")
    note["A1"] = "「%s」怎么填" % spec["title"]
    note["A1"].font = Font(bold=True, size=13)
    note.column_dimensions["A"].width = 96
    r = 3
    note.cell(row=r, column=1, value="示例行请替换成你自己的内容，表头保留。")
    r += 2
    for line in spec["notes"]:
        note.cell(row=r, column=1, value=line).alignment = Alignment(
            wrap_text=False, vertical="top")
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "模板_%s.xlsx" % spec["title"]
