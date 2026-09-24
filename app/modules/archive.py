# -*- coding: utf-8 -*-
"""产出归档:每份周报 / 汇报 / 表格都在飞书留一份,并在「产出记录」多维表格里记一行。

  文档类(周报、汇报)     -> 飞书云文档(调用方已建好,这里只记链接)
  表格类(拓词工作簿等)   -> 飞书在线电子表格(第一张表,点开就看)+ 原始 xlsx 上传云盘(可下载)
  记录                   -> 多维表格一行:时间 / 操作人 / 类型 / 客户 / 标题 / 三个链接 / 花费

权限:文档与表格给**操作人 + 业主**开权限,不对全员开放 —— 「成员只看自己的,业主看全部」。
记录表只给业主;成员在工作台里按自己的 open_id 看自己的记录。

归档失败绝不能拖垮产出本身:所有入口都吞异常、只写日志。
"""
import datetime as dt
import json
import urllib.error
import urllib.request
import uuid

from app import config
from app.modules.social import feishu_docs as fd

FIELDS = [  # (字段名, 类型) 1 文本 2 数字 3 单选 5 日期 15 超链接
    ("标题", 1), ("类型", 3), ("客户", 1), ("操作人", 1), ("时间", 5),
    ("飞书文档", 15), ("在线表格", 15), ("xlsx 下载", 15), ("花费(USD)", 2), ("备注", 1),
]
KINDS = ["拓词总表", "排名汇报", "GSC 周报", "社媒周报", "其他"]


class ArchiveError(RuntimeError):
    pass


def _ok(r, what):
    if r.get("code") != 0:
        raise ArchiveError("%s失败(code %s): %s" % (what, r.get("code"), r.get("msg")))
    return r.get("data") or {}


def owner_id():
    return config.get("social.owner_open_id") or config.get("feishu.receive_id")


def grant(tok, token, typ, open_ids, perm="view", log=None):
    """给若干人开权限。重复开、给自己开都不报错。"""
    for oid in dict.fromkeys(i for i in open_ids if i):
        r = fd._call("POST", "/drive/v1/permissions/%s/members?type=%s&need_notification=false" % (token, typ),
                     tok, {"member_type": "openid", "member_id": oid, "perm": perm})
        if r.get("code") not in (0, 1063003) and log:          # 1063003 = 已是协作者
            log("  授权 %s 给 %s 失败:%s" % (typ, oid[-6:], r.get("msg")))


# ---------------------------------------------------------------- 首次初始化

def ensure_setup(log=None):
    """没有就建:归档文件夹 + 产出记录多维表格。token 写回 config 的 archive 段。"""
    log = log or (lambda m: None)
    folder = config.get("archive.folder_token")
    base = config.get("archive.base_token")
    table = config.get("archive.table_id")
    if folder and base and table:
        return folder, base, table
    tok = fd.token()
    parent = config.get("social.folder_token") or ""
    if not folder:
        d = _ok(fd._call("POST", "/drive/v1/files/create_folder", tok,
                         {"name": "SEO工作台产出", "folder_token": parent}), "建归档文件夹")
        folder = d["token"]
        grant(tok, folder, "folder", [owner_id()], "full_access", log)
        log("已建归档文件夹 %s" % d.get("url", folder))
    if not base:
        d = _ok(fd._call("POST", "/bitable/v1/apps", tok, {"name": "SEO工作台产出记录", "folder_token": folder}),
                "建产出记录表")
        base = d["app"]["app_token"]
        grant(tok, base, "bitable", [owner_id()], "full_access", log)
        tabs = _ok(fd._call("GET", "/bitable/v1/apps/%s/tables" % base, tok), "读数据表")
        table = tabs["items"][0]["table_id"]
        _setup_fields(tok, base, table)
        _drop_blank_rows(tok, base, table)
        log("已建产出记录表 %s" % d["app"].get("url", base))
    _save({"folder_token": folder, "base_token": base, "table_id": table})
    return folder, base, table


def _setup_fields(tok, base, table):
    """新表默认有几个空字段:第一列(主字段)改名「标题」,其余删掉,再按 FIELDS 建。"""
    cur = _ok(fd._call("GET", "/bitable/v1/apps/%s/tables/%s/fields" % (base, table), tok), "读字段")["items"]
    primary = next(f for f in cur if f.get("is_primary"))
    _ok(fd._call("PUT", "/bitable/v1/apps/%s/tables/%s/fields/%s" % (base, table, primary["field_id"]), tok,
                 {"field_name": "标题", "type": 1}), "改主字段")
    for f in cur:
        if not f.get("is_primary"):
            fd._call("DELETE", "/bitable/v1/apps/%s/tables/%s/fields/%s" % (base, table, f["field_id"]), tok)
    for name, typ in FIELDS[1:]:
        body = {"field_name": name, "type": typ}
        if typ == 3:
            body["property"] = {"options": [{"name": k} for k in KINDS]}
        if typ == 5:
            body["property"] = {"date_formatter": "yyyy/MM/dd HH:mm"}
        if typ == 2:
            body["property"] = {"formatter": "0.0000"}
        _ok(fd._call("POST", "/bitable/v1/apps/%s/tables/%s/fields" % (base, table), tok, body), "建字段 " + name)


def _drop_blank_rows(tok, base, table):
    """飞书新建的多维表格自带 10 条空行,不删会混进产出记录。"""
    d = fd._call("GET", "/bitable/v1/apps/%s/tables/%s/records?page_size=100" % (base, table), tok)
    ids = [it["record_id"] for it in ((d.get("data") or {}).get("items") or []) if not (it.get("fields") or {}).get("标题")]
    if ids:
        fd._call("POST", "/bitable/v1/apps/%s/tables/%s/records/batch_delete" % (base, table), tok, {"records": ids})


def operator():
    """当前操作人。上线接飞书登录后由登录态给;本机版读配置,没配就记「本机」。"""
    return config.get("archive.operator") or "本机"


def _save(vals):
    """写回 config.local.yaml 的 archive 段(段落定位,保注释,不碰别的段)。"""
    import re
    p = config.LOCAL
    lines = p.read_text(encoding="utf-8").split("\n") if p.exists() else []
    start = next((i for i, l in enumerate(lines) if re.match(r"^archive\s*:\s*(#.*)?$", l)), None)
    if start is None:
        lines += ["", "# ---------- 产出归档(自动生成,别手改)----------", "archive:"]
        start = len(lines) - 1
    for k, v in vals.items():
        end = next((i for i in range(start + 1, len(lines))
                    if lines[i] and not lines[i].startswith((" ", "\t", "#"))), len(lines))
        hit = next((i for i in range(start + 1, end) if re.match(r"^\s+%s\s*:" % k, lines[i])), None)
        line = '  %s: "%s"' % (k, v)
        if hit is None:
            lines.insert(start + 1, line)
        else:
            lines[hit] = line
    p.write_text("\n".join(lines), encoding="utf-8")
    config.load(refresh=True)


# ---------------------------------------------------------------- 表格

def _upload(tok, path, folder):
    data = path.read_bytes()
    bd = uuid.uuid4().hex
    parts = []
    for k, v in (("file_name", path.name), ("parent_type", "explorer"), ("parent_node", folder), ("size", str(len(data)))):
        parts.append(('--%s\r\nContent-Disposition: form-data; name="%s"\r\n\r\n%s\r\n' % (bd, k, v)).encode("utf-8"))
    parts.append(('--%s\r\nContent-Disposition: form-data; name="file"; filename="%s"\r\n'
                  'Content-Type: application/octet-stream\r\n\r\n' % (bd, path.name)).encode("utf-8") + data + b"\r\n")
    body = b"".join(parts) + ("--%s--\r\n" % bd).encode()
    req = urllib.request.Request(fd.API + "/drive/v1/files/upload_all", data=body, method="POST",
                                 headers={"Authorization": "Bearer " + tok,
                                          "Content-Type": "multipart/form-data; boundary=" + bd})
    try:
        r = json.load(urllib.request.urlopen(req, timeout=120))
    except urllib.error.HTTPError as e:
        r = json.loads(e.read().decode("utf-8", "replace"))
    return _ok(r, "上传 xlsx")["file_token"]


def _col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def publish_table(xlsx_path, title, open_ids, log=None, max_rows=5000):
    """xlsx -> 飞书在线表格(第一张表)+ 原文件上传。返回 (在线表格 url, xlsx url)。"""
    log = log or (lambda m: None)
    from openpyxl import load_workbook
    folder, _, _ = ensure_setup(log)
    tok = fd.token()
    tenant = fd.doc_url("x").rsplit("/docx/", 1)[0]
    ft = _upload(tok, xlsx_path, folder)
    grant(tok, ft, "file", open_ids, "view", log)
    file_url = "%s/file/%s" % (tenant, ft)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)][:max_rows + 1]
    wb.close()
    d = _ok(fd._call("POST", "/sheets/v3/spreadsheets", tok, {"title": title, "folder_token": folder}), "建在线表格")
    st = d["spreadsheet"]["spreadsheet_token"]
    sid = _ok(fd._call("GET", "/sheets/v3/spreadsheets/%s/sheets/query" % st, tok), "读工作表")["sheets"][0]["sheet_id"]
    width = max((len(r) for r in rows), default=1)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]
    for i in range(0, len(rows), 1000):                       # 单次写入有上限,分段
        chunk = rows[i:i + 1000]
        rng = "%s!A%d:%s%d" % (sid, i + 1, _col(width), i + len(chunk))
        _ok(fd._call("PUT", "/sheets/v2/spreadsheets/%s/values" % st, tok,
                     {"valueRange": {"range": rng, "values": chunk}}), "写在线表格")
    grant(tok, st, "sheet", open_ids, "view", log)
    log("在线表格 %d 行已建;xlsx 已上传" % (len(rows) - 1))
    return d["spreadsheet"]["url"], file_url


# ---------------------------------------------------------------- 记录

def record(kind, title, client="", operator="", doc_url=None, sheet_url=None, file_url=None,
           cost=None, note="", log=None):
    log = log or (lambda m: None)
    _, base, table = ensure_setup(log)
    tok = fd.token()
    link = lambda u, t: {"link": u, "text": t} if u else None
    f = {"标题": title, "类型": kind if kind in KINDS else "其他", "客户": client or "",
         "操作人": operator or "", "时间": int(dt.datetime.now().timestamp() * 1000),
         "飞书文档": link(doc_url, "打开文档"), "在线表格": link(sheet_url, "在线查看"),
         "xlsx 下载": link(file_url, "下载 xlsx"), "备注": note or ""}
    if cost is not None:
        f["花费(USD)"] = float(cost)
    f = {k: v for k, v in f.items() if v not in (None, "")}
    d = _ok(fd._call("POST", "/bitable/v1/apps/%s/tables/%s/records" % (base, table), tok, {"fields": f}), "写产出记录")
    log("已记入产出记录表")
    return d["record"]["record_id"]


def list_records(operator=None, limit=100):
    """工作台「我的产出」用。operator 为空 = 全部(业主)。"""
    _, base, table = ensure_setup()
    tok = fd.token()
    body = {"sort": [{"field_name": "时间", "desc": True}]}
    if operator:
        body["filter"] = {"conjunction": "and", "conditions": [
            {"field_name": "操作人", "operator": "is", "value": [operator]}]}
    d = _ok(fd._call("POST", "/bitable/v1/apps/%s/tables/%s/records/search?page_size=%d" % (base, table, limit),
                     tok, body), "读产出记录")
    out = []
    for it in d.get("items") or []:
        f = it.get("fields") or {}
        txt = lambda v: "".join(x.get("text", "") for x in v) if isinstance(v, list) else (v or "")
        url = lambda v: (v or {}).get("link") if isinstance(v, dict) else None
        out.append({"标题": txt(f.get("标题")), "类型": f.get("类型") or "", "客户": txt(f.get("客户")),
                    "操作人": txt(f.get("操作人")), "时间": f.get("时间"),
                    "doc": url(f.get("飞书文档")), "sheet": url(f.get("在线表格")), "file": url(f.get("xlsx 下载")),
                    "花费": f.get("花费(USD)")})
    return out


def safe(fn, *a, log=None, **kw):
    """归档入口统一包一层:失败只写日志,不影响产出本身。"""
    try:
        return fn(*a, log=log, **kw)
    except Exception as e:
        if log:
            log("[归档失败,不影响本次产出] %s: %s" % (type(e).__name__, str(e)[:200]))
        return None
