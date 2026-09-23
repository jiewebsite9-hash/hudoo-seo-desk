# -*- coding: utf-8 -*-
"""上传数据的摄取层:zip / 单文件 -> 一组带路径的工作表。

**为什么不能直接用 server.parse_table_bytes**:那个函数是给词表用的,只读第一个 sheet、
认不出 .xls、也不处理压缩包。社媒后台导出这三样全占:
  · LinkedIn 导出的是 **真 OLE2 .xls**(不是改名的 xlsx),而且**一个文件 6~7 个 sheet**
    (数据 / 新关注者 / 地点 / 职能类别 / 高级 / 所属行业 / 公司规模),只读第一个会丢掉全部受众画像
  · Facebook 导出的是 **UTF-16 CSV**,前三行是 `sep=,` + 指标名 + 表头,第四行才是数据
  · 专员实际交付的是一个压缩包,而且多半是 Mac 打的

**Mac zip 的文件名乱码**:macOS 归档工具把文件名按 UTF-8 存进 zip,但**不置 UTF-8 标志位**
(general purpose bit 11 / EFS)。Python 的 zipfile 见标志位没开就按 cp437 解码,
Windows 资源管理器则按 ANSI 代码页(中文机是 GBK)解码 —— 两边都出乱码。
还原办法是把解错的名字按原编码 encode 回字节再按 UTF-8 decode,见 `_fix_zip_name`。
"""
import io
import re
import zipfile

MAX_UPLOAD = 60 * 1024 * 1024        # 整包上限
MAX_MEMBER = 20 * 1024 * 1024        # 单个成员文件上限
TEXT_ENCODINGS = ("utf-8-sig", "utf-16", "utf-8", "gb18030", "big5")

# Mac 归档和 Windows 缩略图留下的垃圾,一律不当数据文件
JUNK = re.compile(r"(^|/)(__MACOSX/|\.DS_Store$|~\$|\._)")
DATA_EXT = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls")


class IngestError(RuntimeError):
    pass


class Sheet(object):
    """一张解析出来的工作表。

    path   压缩包内的相对路径(已还原中文),单文件上传时就是文件名
    sheet  工作表名;csv 类文件为 ""
    rows   [[单元格字符串, ...], ...],**不剥表头**
           —— 剥不剥、剥几行由各平台解析器决定,这里只负责忠实还原
    note   识别说明,出问题时让人一眼看出是不是认错了格式
    """

    def __init__(self, path, sheet, rows, note):
        self.path = path
        self.sheet = sheet
        self.rows = rows
        self.note = note

    @property
    def parts(self):
        """路径按 / 切开,用来推客户名和平台名。"""
        return [p for p in self.path.replace("\\", "/").split("/") if p]

    def __repr__(self):
        return "<Sheet %s%s %d行>" % (self.path,
                                      ("#" + self.sheet) if self.sheet else "", len(self.rows))


def _fix_zip_name(info):
    """还原 zip 成员的真实文件名。

    zipfile 在 EFS 标志位没置时用 cp437 解码,中文名会变成一串拉丁字符;
    按 cp437 编回字节再按 UTF-8 解,就能拿回原名。编不回去说明本来就是 ASCII 或
    已经是正确的 UTF-8(标志位有置),原样返回。
    """
    name = info.filename
    if info.flag_bits & 0x800:      # EFS 已置位,zipfile 已按 UTF-8 解好
        return name
    try:
        return name.encode("cp437").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name


def read_upload(raw, filename):
    """入口:一个上传的字节流 -> [Sheet, ...]。

    zip 会被展开;单文件按自身格式解析。返回的 Sheet 顺序即压缩包内顺序。
    """
    if len(raw) > MAX_UPLOAD:
        raise IngestError("整包超过 %d MB" % (MAX_UPLOAD // 1024 // 1024))
    if filename.lower().endswith(".zip") or raw[:2] == b"PK":
        return _read_zip(raw)
    return parse_member(raw, filename)


def _read_zip(raw):
    out = []
    skipped = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as e:
        raise IngestError("压缩包打不开:%s" % e)
    with zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = _fix_zip_name(info)
            if JUNK.search("/" + name):
                continue
            if not name.lower().endswith(DATA_EXT):
                skipped.append(name)
                continue
            if info.file_size > MAX_MEMBER:
                skipped.append("%s(超过 %dMB)" % (name, MAX_MEMBER // 1024 // 1024))
                continue
            try:
                out.extend(parse_member(zf.read(info), name))
            except IngestError as e:
                skipped.append("%s(%s)" % (name, e))
    if not out:
        raise IngestError("压缩包里没找到能解析的数据文件"
                          + ("；跳过了:%s" % "、".join(skipped[:5]) if skipped else ""))
    return out


def parse_member(raw, path):
    """单个文件 -> [Sheet, ...]。xls/xlsx 一个 sheet 一条。"""
    low = path.lower()
    if low.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(raw, path)
    if low.endswith(".xls"):
        return _read_xls(raw, path)
    return _read_text(raw, path)


def _read_xlsx(raw, path):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = []
    try:
        for name in wb.sheetnames:
            rows = _tidy([[c for c in r] for r in wb[name].iter_rows(values_only=True)])
            if rows:
                out.append(Sheet(path, name, rows, "xlsx"))
    finally:
        wb.close()
    return out


def _read_xls(raw, path):
    """老式 OLE2 .xls —— LinkedIn 的三个导出文件都是这个格式。

    openpyxl 完全不认它(会报 InvalidFileException),必须用 xlrd。
    xlrd 2.x 砍掉了 xlsx 支持但保留 xls,正好够用。
    """
    try:
        import xlrd
    except ImportError:
        raise IngestError("缺少 xlrd,装一下:pip install xlrd")
    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as e:
        raise IngestError("xls 打不开:%s" % e)

    def cell(ws, r, c):
        """日期单元格必须显式转换 —— xlrd 的 cell_value 对日期返回的是
        Excel 序列号浮点(2026-09-12 会变成 46277.0),直接用就把日期丢了。"""
        cl = ws.cell(r, c)
        if cl.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(cl.value, book.datemode)
            except Exception:
                return cl.value
        return cl.value

    out = []
    for ws in book.sheets():
        rows = _tidy([[cell(ws, r, c) for c in range(ws.ncols)]
                      for r in range(ws.nrows)])
        if rows:
            out.append(Sheet(path, ws.name, rows, "xls(OLE2)"))
    return out


def _read_text(raw, path):
    """csv / tsv / txt。编码逐个试,UTF-16 放前面 —— Facebook 导出就是它。"""
    text = enc = None
    for cand in TEXT_ENCODINGS:
        try:
            t = raw.decode(cand)
        except (UnicodeDecodeError, UnicodeError):
            continue
        if "\x00" in t:          # UTF-16 认错时的典型表现是夹杂大量 NUL
            continue
        text, enc = t, cand
        break
    if text is None:
        raise IngestError("认不出文件编码")

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")

    # Facebook 首行是 `sep=,`,既声明了分隔符也不是数据,读掉
    sep = None
    if lines and lines[0].lower().startswith("sep="):
        sep = lines[0][4:].strip() or ","
        lines = lines[1:]
    if sep is None:
        sample = "\n".join(lines[:40])
        sep = "\t" if ("\t" in sample and sample.count("\t") >= sample.count(",")) else ","

    import csv as _csv
    rows = _tidy(list(_csv.reader(lines, delimiter=sep)))
    if not rows:
        raise IngestError("文件里没读到内容")
    return [Sheet(path, "", rows, "%s · %s 分隔" % (enc, "Tab" if sep == "\t" else sep))]


def _tidy(rows):
    """统一成字符串矩阵:去掉全空行、去掉右侧空列、数字转成紧凑写法。

    Excel 读出来的整数会带 .0(1234.0),直接 str() 会让后面的数字比对和展示都变难看,
    这里统一收干净。日期对象保留 ISO 写法,各平台解析器自己再规整。
    """
    import datetime as _dt
    out = []
    for r in rows:
        cells = []
        for c in r:
            if c is None:
                cells.append("")
            elif isinstance(c, bool):
                cells.append("1" if c else "0")
            elif isinstance(c, float) and c == int(c) and abs(c) < 1e15:
                cells.append(str(int(c)))
            elif isinstance(c, (_dt.datetime, _dt.date)):
                cells.append(c.isoformat()[:19])
            else:
                cells.append(str(c).strip())
        while cells and not cells[-1]:
            cells.pop()
        if cells and any(cells):
            out.append(cells)
    return out
