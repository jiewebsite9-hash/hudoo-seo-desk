# -*- coding: utf-8 -*-
"""步骤 6：汇总基线 summary.json + 与上一次审计对比 + 生成交付 xlsx

第一页是「逐页明细」(在线表格取的就是第一页);问题清单页来自 findings.json(规则引擎初稿)。
"""
import collections, glob, json, os, re, statistics
from urllib.parse import urlparse

from app.modules.audit import findings as F
from app.modules.audit.common import jload, jdump
from app.modules.audit.score import thresholds


def build_summary(cfg):
    out = cfg["out"]
    P, S = jload(out, "pages.json", {}), jload(out, "scores.json", {})
    probe, lh = jload(out, "probe.json", {}), jload(out, "lighthouse.json", {})
    rows = [json.loads(l) for l in open(os.path.join(out, "crawl.jsonl"), encoding="utf-8")]
    sm = jload(out, "sitemap_urls.json", [])
    imgs = jload(out, "images_head.json", [])
    th = thresholds(cfg)
    m = collections.OrderedDict()

    def put(k, v):
        m[k] = v

    st = collections.Counter(r["status"] for r in rows)
    put("抓取 URL 数", len(rows))
    put("HTML 200 页数", len(P))
    for lang, c in sorted(collections.Counter(d["lang"] for d in P.values()).items()):
        put("页面数-%s" % lang, c)
    put("3xx 数", sum(c for s, c in st.items() if 300 <= s < 400))
    put("4xx 数", sum(c for s, c in st.items() if 400 <= s < 500))
    put("5xx 数", sum(c for s, c in st.items() if s >= 500))
    put("请求失败数", st.get(-1, 0))
    put("sitemap URL 数", len(sm))
    put("sitemap 内各语种", dict(collections.Counter(P[u]["lang"] for u in P if P[u]["in_sitemap"])))
    put("可抓页面不在 sitemap", sum(1 for d in P.values() if not d["in_sitemap"]))
    put("noindex 页", sum(1 for d in P.values() if "noindex" in (d["robots_meta"] + d["x_robots"]).lower()))
    put("缺 canonical", sum(1 for d in P.values() if not d["canonical"]))
    put("canonical 相对路径", sum(1 for d in P.values() if d["canonical"] and not d["canonical"].startswith("http")))
    put("canonical 指向他页", sum(1 for d in P.values() if d["canonical"].startswith("http") and d["canonical"].rstrip("/") != d["url"].rstrip("/")))
    put("有 hreflang 的页", sum(1 for d in P.values() if d["hreflang"]))
    put("hreflang 含 x-default 的页", sum(1 for d in P.values() if any(l.lower() == "x-default" for l, _ in d["hreflang"])))
    put("有 JSON-LD 的页", sum(1 for d in P.values() if d["schema"]))
    put("Schema 类型分布", dict(collections.Counter(t for d in P.values() for t in d["schema"]).most_common(12)))
    put("有 og:image 的页", sum(1 for d in P.values() if d["og"].get("image")))
    put("缺 Title", sum(1 for d in P.values() if not d["title"]))
    put("Title 重复页", sum(1 for d in P.values() if d["dup_title"]))
    put("Title 超 60 字符", sum(1 for d in P.values() if len(d["title"]) > 60))
    put("缺 Description", sum(1 for d in P.values() if not d["desc"]))
    put("Description 重复页", sum(1 for d in P.values() if d["dup_desc"]))
    put("缺 H1", sum(1 for d in P.values() if not d["h1"]))
    put("多个 H1", sum(1 for d in P.values() if len(d["h1"]) > 1))
    put("混合内容引用(http://)页", sum(1 for d in P.values() if d["mixed"]))
    put("最大点击深度", max([d["depth"] for d in P.values()] or [0]))
    put("深度>3 的页", sum(1 for d in P.values() if d["depth"] > 3))
    put("正文 0 入链页(非首页)", sum(1 for d in P.values() if d["inlinks"] == 0 and d["type"] != "home"))
    put("断链/畸形链接", len(jload(out, "odd_links.json", {})))

    bi = [i for d in P.values() for i in d["imgs"] if not i["tpl"] and i["src"]]
    put("正文图(全站引用次数)", len(bi))
    put("正文图 WebP/AVIF", sum(1 for i in bi if re.search(r"\.(webp|avif)(\?|$)", i["src"], re.I)))
    put("正文图无 width/height", sum(1 for i in bi if not (i["w"] and i["h"])))
    put("正文图懒加载", sum(1 for i in bi if i["lazy"]))
    put("正文图缺 ALT", sum(1 for i in bi if not (i["alt"] or "").strip()))
    ok = [i for i in imgs if i[1] == 200]
    put("唯一正文图(主语种)", len(ok))
    put("唯一正文图合计 MB", round(sum(i[2] for i in ok) / 1048576.0, 1))
    put("正文图 >%dKB" % th["img_kb_max"], sum(1 for i in ok if i[2] > th["img_kb_max"] * 1024))

    for lang in sorted({v["lang"] for v in S.values()}):
        sc = [v["score"] for v in S.values() if v["lang"] == lang]
        put("Onpage 平均分-%s" % lang, round(sum(sc) / len(sc), 1))
    allsc = [v["score"] for v in S.values()]
    put("Onpage 平均分-全站", round(sum(allsc) / len(allsc), 1) if allsc else None)

    put("CMS", probe.get("cms"))
    put("缺失安全响应头", probe.get("security_missing"))
    put("首页 Cache-Control", probe.get("home_headers", {}).get("cache-control"))
    put("首页 cf-cache-status", probe.get("home_headers", {}).get("cf-cache-status"))
    for u, ts in probe.get("ttfb", {}).items():
        good = [t for t in ts if t > 0]
        put("TTFB 中位 ms | " + urlparse(u).path, statistics.median(good) if good else None)
    for u, r in lh.items():
        p = urlparse(u).path
        for k in ("performance", "seo", "LCP_ms", "CLS", "TBT_ms"):
            put("LH %s | %s" % (k, p), r.get(k))
    jdump(out, "summary.json", m)
    return m


def previous_summary(cfg):
    here = os.path.normcase(os.path.abspath(cfg["out"]))
    dirs = sorted(d for d in glob.glob(os.path.join(os.path.dirname(cfg["out"]), "audit-*"))
                  if os.path.normcase(os.path.abspath(d)) != here and os.path.exists(os.path.join(d, "summary.json")))
    dirs = [d for d in dirs if os.path.basename(d) < os.path.basename(cfg["out"])]
    if not dirs:
        return None, None
    return os.path.basename(dirs[-1]), jload(dirs[-1], "summary.json")


def write_xlsx(cfg, summary, prev_name, prev):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    out = cfg["out"]
    P, S = jload(out, "pages.json", {}), jload(out, "scores.json", {})
    probe, lh = jload(out, "probe.json", {}), jload(out, "lighthouse.json", {})
    issues = jload(out, "findings.json", [])
    th = thresholds(cfg)
    HF, HB = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    RED, YEL, GRN = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "C6EFCE"))
    wb = Workbook()

    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    class Sheet:
        """抓回来的文本可能带控制字符(异常编码、二进制响应),openpyxl 遇到直接抛错 —— 统一先洗"""
        def __init__(self, ws):
            self.ws = ws

        def append(self, row):
            self.ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])

        def __getattr__(self, k):
            return getattr(self.ws, k)

        def __getitem__(self, k):
            return self.ws[k]

    def sheet(title, cols, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
            ws.cell(1, i).font, ws.cell(1, i).fill = HF, HB
        ws.freeze_panes = "A2"
        return Sheet(ws)

    head, rows = F.detail_rows(cfg, F.context(cfg))
    ws = sheet("逐页明细", head, [8, 50, 6, 7, 50, 8, 8, 6, 36, 10, 26, 12, 9, 9, 8, 9, 8, 70], first=True)
    for r in rows:
        ws.append(r)

    ws = sheet("00-基线总览", ["指标", "本次", "上次(%s)" % (prev_name or "无"), "变化"], [34, 40, 40, 14])
    for k, v in summary.items():
        pv = (prev or {}).get(k)
        diff = round(v - pv, 1) if isinstance(v, (int, float)) and isinstance(pv, (int, float)) else ""
        ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v,
                   json.dumps(pv, ensure_ascii=False) if isinstance(pv, (dict, list)) else pv, diff])

    if issues:
        ws = sheet("01-问题清单P0P1P2", ["#", "级别", "维度", "问题", "影响页数", "证据", "动作", "示例页", "负责人", "状态"],
                   [7, 6, 10, 40, 8, 50, 50, 60, 10, 10])
        for it in issues:
            ws.append([it["id"], it["level"], it["dim"], it["title"], it["count"], it["evidence"], it["action"],
                       "\n".join(it["examples"]), "", "待办"])
            ws.cell(ws.max_row, 2).fill = {"P0": RED, "P1": YEL}.get(it["level"], GRN)

    dl = cfg["default_lang"]
    main = [v for v in S.values() if v["lang"] == dl] or list(S.values())
    ws = sheet("02-Onpage站级(%s)" % dl, ["维度", "检查项", "达标页", "失分页", "达标率"], [14, 44, 10, 10, 10])
    agg = collections.OrderedDict()
    for v in main:
        for dim, name, s, _ in v["items"]:
            a = agg.setdefault((dim, name), [0, 0])
            a[0 if s else 1] += 1
    for (dim, name), (g, b) in agg.items():
        rate = g / max(g + b, 1)
        ws.append([dim, name, g, b, "%.0f%%" % (rate * 100)])
        ws.cell(ws.max_row, 5).fill = GRN if rate >= 0.8 else (YEL if rate >= 0.5 else RED)

    names = [n for _, n, _, _ in next(iter(S.values()))["items"]] if S else []
    ws = sheet("03-Onpage逐页明细", ["URL", "语种", "页型", "推断主词", "得分", "评级"] + names, [55, 6, 16, 22, 6, 9] + [6] * len(names))
    for u, v in sorted(S.items(), key=lambda x: (x[1]["lang"], x[1]["score"])):
        ws.append([u, v["lang"], v["type"], v["kw"], v["score"], v["rating"]] + ["✅" if s else "❌" for _, _, s, _ in v["items"]])
        ws.cell(ws.max_row, 6).fill = GRN if v["score"] >= th["good"] else (YEL if v["score"] >= th["ok"] else RED)
    for c in range(7, 7 + len(names)):
        ws.cell(1, c).alignment = Alignment(text_rotation=90, wrap_text=True)
    ws.row_dimensions[1].height = 170

    ws = sheet("05-技术实测", ["类别", "URL / 项", "结果", "说明"], [14, 60, 40, 60])
    for u, ch in probe.get("redirects", {}).items():
        ws.append(["主机归一", u, " → ".join(str(h.get("status", h.get("error", ""))) for h in ch), ch[-1].get("url", "") if ch else ""])
    for u, ch in probe.get("not_found", {}).items():
        last = ch[-1] if ch else {}
        ws.append(["404 行为", u, last.get("status"), "软404风险" if last.get("status") == 200 else last.get("title", "")])
    for p, r in probe.get("shadow", {}).items():
        ws.append(["影子URL(%s)" % probe.get("cms"), p, r.get("final_status"), r.get("title", "")])
    ws.append(["响应头", "缺失安全头", ", ".join(probe.get("security_missing", [])) or "无", ""])
    for k, v in probe.get("home_headers", {}).items():
        ws.append(["响应头", "首页 " + k, v, ""])
    for u, r in probe.get("asset_cache", {}).items():
        ws.append(["静态资源缓存", u, r.get("cc", r.get("error", "")), "cf=%s enc=%s" % (r.get("cf", ""), r.get("enc", ""))])
    for u, ts in probe.get("ttfb", {}).items():
        ws.append(["TTFB×5(ms)", u, str(ts), "标准 ≤%dms" % th["ttfb_ms"]])
    for u, r in lh.items():
        ws.append(["Lighthouse 移动", u, "性能 %s / SEO %s" % (r.get("performance"), r.get("seo")),
                   "LCP %sms(≤%s) CLS %s(≤%s) TBT %sms；SEO未过：%s" % (r.get("LCP_ms"), th["lcp_ms"], r.get("CLS"), th["cls"], r.get("TBT_ms"),
                                                                     "；".join(r.get("seo_fails", [])) or "无")])
    ws.append(["robots.txt", "状态 %s" % probe.get("robots", {}).get("status"), "", (probe.get("robots", {}).get("text") or "")[:1000]])
    for c in ws["D"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = sheet("06-图片体积", ["图片 URL", "状态", "KB", "类型", "Cache-Control"], [80, 6, 10, 16, 30])
    for i in sorted(jload(out, "images_head.json", []), key=lambda x: -x[2]):
        ws.append([i[0], i[1], round(i[2] / 1024.0), i[3], i[4]])
        if i[2] > th["img_kb_max"] * 1024:
            ws.cell(ws.max_row, 3).fill = RED
        elif i[2] > th["img_kb_ideal"] * 1024:
            ws.cell(ws.max_row, 3).fill = YEL

    ws = sheet("07-断链", ["目标 [状态]", "出现在哪些页"], [70, 90])
    for k, v in jload(out, "odd_links.json", {}).items():
        ws.append([k, "\n".join(v[:20])])

    ws = sheet("08-口径说明", ["项", "说明"], [26, 110])
    for row in [
        ("审计依据", "On-page 6 维 30 项 + 技术审计 7 维；阈值见下行（来自工作台配置 audit.thresholds）"),
        ("本次阈值", json.dumps(th, ensure_ascii=False)),
        ("抓取", "sitemap（含索引嵌套）+ 同域 BFS，上限 %s 页；UA=桌面 Chrome；数据时点 %s" % (cfg.get("max_pages"), cfg.get("date"))),
        ("Onpage 评级", "≥%d 良好 / %d-%d 需优化 / <%d 全面整改；修复后须复测" % (th["good"], th["ok"], th["good"] - 1, th["ok"])),
        ("关键词密度", "不作硬性目标，不计分"),
        ("主关键词", "从 Title/H1 剥品牌词推断，可能不准；以拓词后的关键词-URL 映射为准"),
        ("代理指标", "「段落≤4句」用正文<1200词近似；「内容无过时信息」「外链 nofollow 合理」脚本不判、默认通过，须人工复核"),
        ("正文判定", "剥 header/footer/nav 及常见头尾 class；出现在 ≥80% 页面的链接/图片视为模板件，不计入正文"),
        ("Lighthouse", "移动端实验室数据，非 CrUX 实况；INP 须看 CrUX/GSC，TBT 仅作代理"),
        ("图片体积", "仅主语种正文图 HEAD 实测，上限 800 张"),
    ]:
        ws.append(list(row))
    path = os.path.join(out, "%s-SEO审计表-%s.xlsx" % (cfg["site"], cfg["date"]))
    wb.save(path)
    return path


def run(cfg, log=print):
    summary = build_summary(cfg)
    prev_name, prev = previous_summary(cfg)
    path = write_xlsx(cfg, summary, prev_name, prev)
    log("summary.json 已写；对比基线：%s" % (prev_name or "无（首次审计）"))
    log("xlsx：" + os.path.basename(path))
